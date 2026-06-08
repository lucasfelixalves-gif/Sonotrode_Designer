import os
import time
from pathlib import Path

import pandas as pd
from PySide6.QtCore import Qt
from PySide6.QtGui import QCursor, QTextCursor
from PySide6.QtWidgets import QAbstractItemView, QInputDialog, QMessageBox, QToolTip

from .base_controller import BaseController

try:
    from scripts.ccx_solver import CCXSolverWorker
    from scripts.app.workers.mesh_worker import MeshGenerationWorker
    from scripts.analytics_service import load_geometry_parameters
    from scripts.augment_master_workbook import create_master_workbook
    from scripts.analytics_table_model import DataFrameTableModel
    from scripts.ccx_results_parser import (
        SUMMARY_COLUMNS,
        build_summary_row,
        parse_dat_modal_results,
        parse_mesh_node_sets,
        select_longitudinal_mode,
    )
except ModuleNotFoundError:
    from ccx_solver import CCXSolverWorker
    from app.workers.mesh_worker import MeshGenerationWorker
    from analytics_service import load_geometry_parameters
    try:
        from augment_master_workbook import create_master_workbook
    except Exception:
        create_master_workbook = None
    from analytics_table_model import DataFrameTableModel
    from ccx_results_parser import (
        SUMMARY_COLUMNS,
        build_summary_row,
        parse_dat_modal_results,
        parse_mesh_node_sets,
        select_longitudinal_mode,
    )


class AnalyticsController(BaseController):
    """Controller for Tab 2: job manager and analytics dashboard."""

    DEFAULT_RESULTS_SUMMARY_FILENAME = "Results_Summary.xlsx"

    CORE_TABLE_COLUMNS = [
        "Model_name",
        "Status",
        "Number_of_nodes",
        "Mode_Index",
        "Frequency_Hz",
        "YMAX_Mean_µm",
        "YMIN_Mean_µm",
        "Gain_Ratio",
        "Tip_Std_µm",
        "Base_Std_µm",
    ]

    def _disconnect_plot_hover(self):
        hover_cid = getattr(self, "_plot_hover_cid", None)
        if hover_cid is not None:
            try:
                self.widgets.canvas.mpl_disconnect(hover_cid)
            except Exception:  # noqa: BLE001
                pass
        self._plot_hover_cid = None
        self._plot_hover_artist = None
        self._plot_hover_points = []
        QToolTip.hideText()

    def _bind_plot_hover(self, scatter_artist, points):
        self._disconnect_plot_hover()
        self._plot_hover_artist = scatter_artist
        self._plot_hover_points = points
        self._plot_hover_cid = self.widgets.canvas.mpl_connect("motion_notify_event", self._on_plot_hover)

    def _on_plot_hover(self, event):
        artist = getattr(self, "_plot_hover_artist", None)
        points = getattr(self, "_plot_hover_points", [])
        if artist is None or not points or event is None or event.inaxes is None:
            QToolTip.hideText()
            return

        try:
            contains, details = artist.contains(event)
        except Exception:  # noqa: BLE001
            QToolTip.hideText()
            return

        if not contains:
            QToolTip.hideText()
            return

        indices = details.get("ind", []) if isinstance(details, dict) else []
        if not indices:
            QToolTip.hideText()
            return

        idx = int(indices[0])
        if idx < 0 or idx >= len(points):
            QToolTip.hideText()
            return

        point = points[idx]
        model_name = point.get("model_name", "")
        x_val = point.get("x")
        y_val = point.get("y")
        z_val = point.get("z")

        if z_val is None:
            tooltip = f"{model_name} ({x_val:.6g};{y_val:.6g})"
        else:
            tooltip = f"{model_name} ({x_val:.6g};{y_val:.6g};{z_val:.6g})"

        QToolTip.showText(QCursor.pos(), tooltip, self.widgets.canvas)

    def _normalize_model_frame(self, dataframe):
        if dataframe is None:
            return pd.DataFrame()

        frame = dataframe.copy()
        if "Model_name" not in frame.columns:
            return frame.reset_index(drop=True)

        if frame.empty:
            return frame.reset_index(drop=True)

        frame = frame.dropna(subset=["Model_name"])
        frame["Model_name"] = frame["Model_name"].astype(str).str.strip()
        frame = frame[frame["Model_name"].str.lower() != "nan"]
        frame = frame[frame["Model_name"] != ""]
        frame = frame.drop_duplicates(subset=["Model_name"], keep="last")
        return frame.reset_index(drop=True)

    def _find_step_models(self):
        if not self.session.project_dir:
            return []

        step_dir = Path(self.session.project_dir) / "02_Geometry_STEP"
        if not step_dir.is_dir():
            return []

        patterns = ("*.step", "*.STEP", "*.stp", "*.STP")
        step_files = []
        for pattern in patterns:
            step_files.extend(step_dir.glob(pattern))

        models = []
        seen = set()
        for step_path in sorted(step_files, key=lambda path: path.name.lower()):
            model_name = step_path.stem.strip()
            if not model_name or model_name in seen:
                continue
            seen.add(model_name)
            models.append(model_name)
        return models

    def _build_conditional_material_frame(self, model_names):
        return self._build_material_frame(model_names, require_exact_match=True)

    def _has_any_exact_material_match(self, model_names):
        material_df = self._normalize_model_frame(self.session.df_material)
        if material_df.empty or "Model_name" not in material_df.columns:
            return False

        exact_rows = material_df[material_df["Model_name"].astype(str).str.strip() != "*"]
        if exact_rows.empty:
            return False

        project_models = {str(name).strip() for name in model_names if str(name).strip()}
        exact_models = set(exact_rows["Model_name"].astype(str).str.strip())
        return bool(project_models.intersection(exact_models))

    def _build_material_frame(self, model_names, require_exact_match):
        material_df = self._normalize_model_frame(self.session.df_material)
        if material_df.empty or "Model_name" not in material_df.columns:
            return pd.DataFrame()

        material_df["Model_name"] = material_df["Model_name"].astype(str).str.strip()
        material_df = material_df[material_df["Model_name"] != ""]

        exact_rows = material_df[material_df["Model_name"] != "*"]
        fallback_rows = material_df[material_df["Model_name"] == "*"]
        if require_exact_match and not self._has_any_exact_material_match(model_names):
            return pd.DataFrame()

        exact_lookup = exact_rows.drop_duplicates(subset=["Model_name"], keep="last").set_index("Model_name") if not exact_rows.empty else pd.DataFrame()
        fallback_row = fallback_rows.iloc[-1] if not fallback_rows.empty else None
        material_columns = [column for column in material_df.columns if column != "Model_name"]

        rows = []
        for model_name in model_names:
            if not exact_lookup.empty and model_name in exact_lookup.index:
                source_row = exact_lookup.loc[model_name]
            elif fallback_row is not None:
                source_row = fallback_row
            else:
                source_row = None

            row = {"Model_name": model_name}
            for column in material_columns:
                row[column] = source_row.get(column, pd.NA) if source_row is not None else pd.NA
            rows.append(row)

        return pd.DataFrame(rows, columns=["Model_name", *material_columns])

    def _build_full_merged_dataframe(self, results_df):
        merged = self._normalize_model_frame(results_df)
        if merged.empty:
            return merged

        model_names = merged["Model_name"].astype(str).str.strip().tolist()

        geometry_df = pd.DataFrame()
        try:
            geometry_df = load_geometry_parameters(self.session.project_dir)
        except Exception as exc:  # noqa: BLE001
            print(f"[WARNING] Could not load geometry parameters from master workbook: {exc}")

        geometry_df = self._normalize_model_frame(geometry_df)
        if not geometry_df.empty:
            geometry_df = geometry_df[geometry_df["Model_name"].isin(model_names)].copy()
            overlapping_columns = [
                column for column in geometry_df.columns
                if column != "Model_name" and column in merged.columns
            ]
            if overlapping_columns:
                geometry_df = geometry_df.drop(columns=overlapping_columns)
            merged = merged.merge(geometry_df, on="Model_name", how="left")

        material_df = self._build_material_frame(model_names, require_exact_match=False)
        if not material_df.empty:
            overlapping_columns = [
                column for column in material_df.columns
                if column != "Model_name" and column in merged.columns
            ]
            if overlapping_columns:
                material_df = material_df.drop(columns=overlapping_columns)
            merged = merged.merge(material_df, on="Model_name", how="left")

        return merged

    def _build_ui_table_dataframe(self, merged_df):
        merged = self._normalize_model_frame(merged_df)
        if merged.empty:
            return merged

        model_count = len(merged.index)
        model_names = merged["Model_name"].astype(str).str.strip().tolist() if "Model_name" in merged.columns else []
        include_material_columns = self._has_any_exact_material_match(model_names)
        material_columns = []
        if self.session.df_material is not None and not self.session.df_material.empty and "Model_name" in self.session.df_material.columns:
            material_columns = [column for column in self.session.df_material.columns if column != "Model_name"]

        ordered_columns = [column for column in self.CORE_TABLE_COLUMNS if column in merged.columns]

        if model_count <= 1:
            varying_columns = [
                column for column in merged.columns
                if column not in ordered_columns and (include_material_columns or column not in material_columns)
            ]
        else:
            varying_columns = [
                column
                for column in merged.columns
                if column not in ordered_columns
                and (include_material_columns or column not in material_columns)
                and merged[column].nunique(dropna=True) > 1
            ]

        ordered_columns.extend(varying_columns)
        return merged.loc[:, [column for column in ordered_columns if column in merged.columns]].copy()

    def _build_export_dataframe(self, merged_df, ui_df=None):
        merged = self._normalize_model_frame(merged_df)
        if merged.empty:
            return merged

        ordered_columns = []
        if ui_df is not None and not ui_df.empty:
            ordered_columns.extend([column for column in ui_df.columns if column in merged.columns])
        else:
            ordered_columns.extend([column for column in self.CORE_TABLE_COLUMNS if column in merged.columns])

        geometry_columns = []
        try:
            geometry_df = load_geometry_parameters(self.session.project_dir)
            if geometry_df is not None and not geometry_df.empty and "Model_name" in geometry_df.columns:
                geometry_columns = [
                    column for column in geometry_df.columns
                    if column != "Model_name" and column in merged.columns
                ]
        except Exception as exc:  # noqa: BLE001
            print(f"[WARNING] Could not load geometry parameters for export ordering: {exc}")

        for column in geometry_columns:
            if column not in ordered_columns:
                ordered_columns.append(column)

        material_columns = []
        if self.session.df_material is not None and not self.session.df_material.empty and "Model_name" in self.session.df_material.columns:
            material_columns = [
                column for column in self.session.df_material.columns
                if column != "Model_name" and column in merged.columns
            ]
        for column in material_columns:
            if column not in ordered_columns:
                ordered_columns.append(column)

        return merged.loc[:, [column for column in ordered_columns if column in merged.columns]].copy()

    def _build_results_rows_from_disk(self):
        table_columns = ["Model_name", "Status", *[column for column in SUMMARY_COLUMNS if column != "Model_name"]]
        if not self.session.project_dir:
            return pd.DataFrame(columns=table_columns)

        mesh_dir = Path(self.session.project_dir) / "03_Meshes_INP"
        results_dir = Path(self.session.project_dir) / "04_Results_CCX"
        rows = []

        for model_name in self._find_step_models():
            mesh_inp_path = mesh_dir / f"{model_name}_mesh.inp"
            dat_path = results_dir / f"{model_name}.dat"
            frd_path = results_dir / f"{model_name}.frd"

            has_mesh = mesh_inp_path.is_file()
            has_solved_result = frd_path.is_file()
            status = "Solved" if has_mesh and has_solved_result else "Meshed" if has_mesh else "Pending"

            row = {
                "Model_name": model_name,
                "Status": status,
                **{column: "--" for column in SUMMARY_COLUMNS if column != "Model_name"},
            }

            if has_mesh:
                try:
                    _, node_count = parse_mesh_node_sets(str(mesh_inp_path))
                    row["Number_of_nodes"] = node_count
                except Exception as exc:  # noqa: BLE001
                    print(f"[WARNING] Could not parse mesh node count for {model_name}: {exc}")

            if status == "Solved" and dat_path.is_file():
                try:
                    node_sets, node_count = parse_mesh_node_sets(str(mesh_inp_path))
                    mode_frequencies, mode_displacements = parse_dat_modal_results(str(dat_path), node_sets)
                    metrics = select_longitudinal_mode(
                        mode_displacements,
                        input_amplitude=self.session.input_amplitude,
                    )
                    if metrics is not None:
                        row.update(build_summary_row(model_name, node_count, mode_frequencies, metrics))
                except Exception as exc:  # noqa: BLE001
                    print(f"[WARNING] Could not parse solver metrics for {model_name}: {exc}")

            rows.append(row)

        return pd.DataFrame(rows, columns=table_columns)

    def _merge_job_table_sources(self, results_df):
        return self._build_full_merged_dataframe(results_df)

    def _numeric_columns_for_plotting(self, dataframe):
        numeric_columns = []
        if dataframe is None or dataframe.empty:
            return numeric_columns

        for column in dataframe.columns:
            if column in {"Model_name", "Status"}:
                continue
            numeric_series = pd.to_numeric(dataframe[column], errors="coerce")
            if numeric_series.notna().any():
                numeric_columns.append(column)
        return numeric_columns

    def _refresh_job_table_from_disk(self):
        self.session.current_model_name = None
        self.session.current_mode_index = None
        self.session.available_modes_dict = {}
        self.session.base_result_mesh = None
        self.session.result_camera_initialized = False
        self.session.selected_result_row = None

        results_df = self._build_results_rows_from_disk()
        self.session.df_results = results_df.copy()
        self.session.df_merged = self._merge_job_table_sources(results_df)
        self.session.job_table_df = self._build_ui_table_dataframe(self.session.df_merged)

        if self.session.table_model is None:
            self.session.table_model = DataFrameTableModel(self.session.job_table_df)
        else:
            self.session.table_model.set_dataframe(self.session.job_table_df)

        if hasattr(self.ui, "table_results"):
            self.ui.table_results.setModel(self.session.table_model)
            self._connect_table_selection_signal()

        if hasattr(self.ui, "label_mode_data"):
            self.ui.label_mode_data.setText("Mode nº: -- | Frequency: -- Hz")
        if hasattr(self.ui, "btn_prev_mode"):
            self.ui.btn_prev_mode.setEnabled(False)
        if hasattr(self.ui, "btn_next_mode"):
            self.ui.btn_next_mode.setEnabled(False)
        if hasattr(self.ui, "btn_render_3D"):
            self.ui.btn_render_3D.setEnabled(False)
        if hasattr(self.ui, "btn_load_3d"):
            self.ui.btn_load_3d.setEnabled(False)
        if hasattr(self.ui, "lbl_info"):
            self.ui.lbl_info.setText("Select a result row to enable 3D loading.")
        if self.widgets.plotter_results is not None:
            self.widgets.plotter_results.clear()

        numeric_columns = self._numeric_columns_for_plotting(self.session.job_table_df)
        if hasattr(self.ui, "combo_x") and hasattr(self.ui, "combo_y") and hasattr(self.ui, "combo_z"):
            self.ui.combo_x.blockSignals(True)
            self.ui.combo_y.blockSignals(True)
            self.ui.combo_z.blockSignals(True)
            self.ui.combo_x.clear()
            self.ui.combo_y.clear()
            self.ui.combo_z.clear()
            self.ui.combo_x.addItems(numeric_columns)
            self.ui.combo_y.addItems(numeric_columns)
            self.ui.combo_z.addItem("--- None (2D) ---")
            self.ui.combo_z.addItems(numeric_columns)
            self.ui.combo_x.blockSignals(False)
            self.ui.combo_y.blockSignals(False)
            self.ui.combo_z.blockSignals(False)

            if len(numeric_columns) >= 2:
                self.ui.combo_x.setCurrentIndex(0)
                self.ui.combo_y.setCurrentIndex(1)
            elif len(numeric_columns) == 1:
                self.ui.combo_x.setCurrentIndex(0)
                self.ui.combo_y.setCurrentIndex(0)
            self.ui.combo_z.setCurrentIndex(0)

        self.update_plot()
        return self.session.job_table_df

    def bind(self):
        if hasattr(self.ui, "combo_x"):
            self.ui.combo_x.currentIndexChanged.connect(self.update_plot)
        if hasattr(self.ui, "combo_y"):
            self.ui.combo_y.currentIndexChanged.connect(self.update_plot)
        if hasattr(self.ui, "combo_z"):
            self.ui.combo_z.currentIndexChanged.connect(self.update_plot)

        if hasattr(self.ui, "table_results"):
            self.ui.table_results.setSelectionBehavior(QAbstractItemView.SelectRows)
            self.ui.table_results.setSelectionMode(QAbstractItemView.ExtendedSelection)
            self._connect_table_selection_signal()

        if hasattr(self.ui, "btn_select_all") and hasattr(self.ui, "table_results"):
            self.ui.btn_select_all.clicked.connect(self.ui.table_results.selectAll)
        if hasattr(self.ui, "btn_clear_selection") and hasattr(self.ui, "table_results"):
            self.ui.btn_clear_selection.clicked.connect(self.ui.table_results.clearSelection)

        if hasattr(self.ui, "btn_generate_meshes"):
            self.ui.btn_generate_meshes.clicked.connect(self.generate_selected_meshes)
        if hasattr(self.ui, "btn_toggle_ccx"):
            self.widgets.btn_toggle_ccx_default_style = self.ui.btn_toggle_ccx.styleSheet()
            self.ui.btn_toggle_ccx.clicked.connect(self.toggle_ccx_solver)
            self._reset_ccx_toggle_button()
        if hasattr(self.ui, "btn_generate_geometry"):
            self.ui.btn_generate_geometry.clicked.connect(self.generate_geometry_workbook)
        if hasattr(self.ui, "btn_export_excel"):
            self.ui.btn_export_excel.clicked.connect(self.export_table_to_excel)

    def _on_table_selection_changed(self, *_):
        self.update_plot()

    def _connect_table_selection_signal(self):
        if not hasattr(self.ui, "table_results"):
            return
        selection_model = self.ui.table_results.selectionModel()
        if selection_model is None:
            return
        try:
            selection_model.selectionChanged.disconnect(self._on_table_selection_changed)
        except Exception:  # noqa: BLE001
            pass
        selection_model.selectionChanged.connect(self._on_table_selection_changed)

    def update_console(self, text):
        if not hasattr(self.ui, "plainTextEdit_console"):
            return
        console = self.ui.plainTextEdit_console
        cursor = console.textCursor()
        cursor.movePosition(QTextCursor.End)
        cursor.insertText(text)
        console.setTextCursor(cursor)
        console.ensureCursorVisible()

    def _set_status_text(self, message):
        if hasattr(self.ui, "lbl_status"):
            self.ui.lbl_status.setText(message)
        elif hasattr(self.ui, "lbl_info"):
            self.ui.lbl_info.setText(message)

    def _set_progress_value(self, value):
        if hasattr(self.ui, "progress_bar"):
            self.ui.progress_bar.setValue(int(value))

    def _log_pipeline_message(self, message):
        self._set_status_text(message)
        if hasattr(self.ui, "plainTextEdit_console"):
            self.update_console(f"{message}\n")

    def _on_mesh_status_updated(self, message):
        parts = str(message or "").split("|", 2)
        if len(parts) >= 3 and parts[0] == "MESH":
            stage = parts[1].strip().upper()
            payload = parts[2].strip()
            model_name = payload.split("(", 1)[0].strip().split(":", 1)[0].strip()
            if model_name:
                if stage == "START":
                    self.update_row_status([model_name], "Meshing...")
                elif stage == "DONE":
                    self.update_row_status([model_name], "Meshed")
                elif stage in {"FAIL", "ERROR"}:
                    self.update_row_status([model_name], "Failed")
        self._log_pipeline_message(message)

    def _on_ccx_status_updated(self, message):
        parts = str(message or "").split("|", 2)
        if len(parts) >= 3 and parts[0] == "SOLVE":
            stage = parts[1].strip().upper()
            payload = parts[2].strip()
            model_name = payload.split("(", 1)[0].strip().split(":", 1)[0].strip()
            if model_name:
                if stage == "START":
                    self.update_row_status([model_name], "Solving...")
                elif stage in {"DONE", "SUMMARY", "COMPLETE"}:
                    self.update_row_status([model_name], "Solved")
                elif stage in {"FAIL", "ERROR"}:
                    self.update_row_status([model_name], "Failed")
                elif stage == "ABORT":
                    self.update_row_status([model_name], "Aborted")
        self._log_pipeline_message(message)

    def update_sim_timer(self):
        if not hasattr(self.ui, "label") and not hasattr(self.ui, "label_sim_timer"):
            return

        if self.session.sim_start_time <= 0:
            if hasattr(self.ui, "label"):
                self.ui.label.setText("ElapsedTime: 00:00")
            if hasattr(self.ui, "label_sim_timer"):
                self.ui.label_sim_timer.setText("Elapsed Time: 00:00")
            return

        elapsed_seconds = max(0, int(time.time() - self.session.sim_start_time))
        minutes, seconds = divmod(elapsed_seconds, 60)
        if hasattr(self.ui, "label"):
            self.ui.label.setText(f"ElapsedTime: {minutes:02d}:{seconds:02d}")
        if hasattr(self.ui, "label_sim_timer"):
            self.ui.label_sim_timer.setText(f"Elapsed Time: {minutes:02d}:{seconds:02d}")

    def get_selected_geometries(self):
        if not hasattr(self.ui, "table_results"):
            return []
        selection_model = self.ui.table_results.selectionModel()
        model = self.ui.table_results.model()
        if selection_model is None or model is None:
            return []

        selected_models = []
        for row_index in selection_model.selectedRows():
            model_index = model.index(row_index.row(), 0)
            model_name = str(model.data(model_index, Qt.DisplayRole) or "").strip()
            if model_name and model_name not in selected_models:
                selected_models.append(model_name)
        return selected_models

    def initialize_job_table(self):
        return self._refresh_job_table_from_disk()

    def update_row_status(self, model_names_list, new_status):
        if self.session.job_table_df is None or self.session.job_table_df.empty:
            return

        target_names = {
            str(name).strip()
            for name in (model_names_list or [])
            if str(name).strip()
        }
        if not target_names or "Model_name" not in self.session.job_table_df.columns:
            return

        mask = self.session.job_table_df["Model_name"].astype(str).str.strip().isin(target_names)
        if not mask.any():
            return

        self.session.job_table_df.loc[mask, "Status"] = str(new_status)
        if self.session.df_merged is not None and not self.session.df_merged.empty and "Model_name" in self.session.df_merged.columns:
            merged_mask = self.session.df_merged["Model_name"].astype(str).str.strip().isin(target_names)
            if merged_mask.any():
                self.session.df_merged.loc[merged_mask, "Status"] = str(new_status)
        if self.session.table_model is not None:
            self.session.table_model.set_dataframe(self.session.job_table_df)
        elif hasattr(self.ui, "table_results"):
            self.session.table_model = DataFrameTableModel(self.session.job_table_df)
            self.ui.table_results.setModel(self.session.table_model)

    def _on_mesh_worker_finished(self):
        self.session.mesh_worker = None
        if hasattr(self.ui, "btn_generate_meshes"):
            self.ui.btn_generate_meshes.setEnabled(True)
        if hasattr(self.ui, "btn_toggle_ccx") and not self.session.is_solving:
            self.ui.btn_toggle_ccx.setEnabled(True)
        self.initialize_job_table()

    def _reset_ccx_toggle_button(self):
        if not hasattr(self.ui, "btn_toggle_ccx"):
            return
        self.ui.btn_toggle_ccx.setText("Run CCX")
        self.ui.btn_toggle_ccx.setStyleSheet(getattr(self.widgets, "btn_toggle_ccx_default_style", ""))

    def _on_ccx_worker_finished(self):
        self.widgets.sim_timer.stop()
        self.session.sim_start_time = 0
        self.update_sim_timer()
        self.session.is_solving = False
        self._reset_ccx_toggle_button()
        if hasattr(self.ui, "btn_tetris"):
            self.ui.btn_tetris.setEnabled(False)

        setup = self.context.controllers.get("setup") if hasattr(self.context, "controllers") else None
        if setup is not None and getattr(setup, "tetris_window", None) is not None and setup.tetris_window.isVisible():
            setup.tetris_window.close()

        if hasattr(self.ui, "btn_generate_meshes"):
            self.ui.btn_generate_meshes.setEnabled(True)
        if hasattr(self.ui, "btn_toggle_ccx") and not (
            self.session.mesh_worker is not None and self.session.mesh_worker.isRunning()
        ):
            self.ui.btn_toggle_ccx.setEnabled(True)

        self.session.ccx_worker = None
        self.initialize_job_table()
        try:
            self.load_analytics_data()
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(
                self.main_window,
                "Analytics Update",
                f"CCX run finished, but analytics refresh failed:\n{exc}",
            )

    def generate_selected_meshes(self):
        selected_models = self.get_selected_geometries()
        if not selected_models:
            QMessageBox.warning(self.main_window, "No Selection", "Select at least one row in the analytics table.")
            return
        if not self.session.project_dir:
            QMessageBox.warning(self.main_window, "No Project", "Load or create a project first.")
            return

        setup = self.context.controllers["setup"]
        try:
            if not setup.ensure_configuration_loaded():
                QMessageBox.warning(
                    self.main_window,
                    "Missing Configuration",
                    "Place exactly one workbook in 01_Master_Config and click Augment or Load.",
                )
                return
        except Exception as exc:  # noqa: BLE001
            QMessageBox.warning(self.main_window, "Missing Configuration", f"Could not load workbook:\n{exc}")
            return

        if self.session.mesh_worker is not None and self.session.mesh_worker.isRunning():
            QMessageBox.information(self.main_window, "Meshing Busy", "Selected meshing is already running.")
            return

        if self.session.is_solving:
            QMessageBox.information(self.main_window, "CCX Running", "Abort CCX before starting a new meshing run.")
            return

        self._set_progress_value(0)
        self._log_pipeline_message("MESH|QUEUE|Starting selected mesh generation")
        if hasattr(self.ui, "btn_toggle_ccx"):
            self.ui.btn_toggle_ccx.setEnabled(False)
        if hasattr(self.ui, "btn_generate_meshes"):
            self.ui.btn_generate_meshes.setEnabled(False)

        self.session.mesh_worker = MeshGenerationWorker(
            project_dir=self.session.project_dir,
            df_mesh=self.session.df_mesh,
            semantic_roles=self.session.semantic_roles,
            target_models=selected_models,
        )

        self.session.mesh_worker.progress_updated.connect(self._set_progress_value)
        self.session.mesh_worker.status_updated.connect(self._on_mesh_status_updated)
        self.session.mesh_worker.finished.connect(self._on_mesh_worker_finished)
        self.session.mesh_worker.start()

    def export_table_to_excel(self):
        if not self.session.project_dir:
            QMessageBox.warning(self.main_window, "No Project", "Load or create a project first.")
            return
        if self.session.df_merged is None or self.session.df_merged.empty:
            QMessageBox.warning(self.main_window, "No Table Data", "There is no job table data to export.")
            return

        reports_dir = os.path.join(self.session.project_dir, "05_Reports")
        os.makedirs(reports_dir, exist_ok=True)

        file_name, ok = QInputDialog.getText(
            self.main_window,
            "Summary Table File Name",
            "Enter the Excel file name for the summary table:",
            text=self.DEFAULT_RESULTS_SUMMARY_FILENAME,
        )
        if not ok:
            return

        file_name = file_name.strip()
        if not file_name:
            QMessageBox.warning(self.main_window, "Invalid File Name", "Please enter a summary file name.")
            return
        if not file_name.lower().endswith(".xlsx"):
            file_name = f"{file_name}.xlsx"

        export_path = os.path.join(reports_dir, file_name)

        export_df = self._build_export_dataframe(self.session.df_merged, self.session.job_table_df)
        export_df.to_excel(export_path, index=False)
        self._format_export_workbook(export_path, export_df.columns)
        self.update_console(f"[EXPORT] Job table exported to: {export_path}\n")

    def _format_export_workbook(self, export_path, export_columns):
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except Exception as exc:  # noqa: BLE001
            self.update_console(f"[EXPORT][WARNING] Could not apply workbook formatting: {exc}\n")
            return

        geometry_columns = []
        try:
            geometry_df = load_geometry_parameters(self.session.project_dir)
            if geometry_df is not None and not geometry_df.empty and "Model_name" in geometry_df.columns:
                geometry_columns = [column for column in geometry_df.columns if column != "Model_name"]
        except Exception as exc:  # noqa: BLE001
            self.update_console(f"[EXPORT][WARNING] Could not load geometry columns for formatting: {exc}\n")

        material_columns = []
        if self.session.df_material is not None and not self.session.df_material.empty and "Model_name" in self.session.df_material.columns:
            material_columns = [column for column in self.session.df_material.columns if column != "Model_name"]

        geometry_set = set(geometry_columns)
        material_set = set(material_columns)

        wb = load_workbook(export_path)
        ws = wb.active

        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        blue_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        yellow_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
        green_fill = PatternFill(start_color="93C47D", end_color="93C47D", fill_type="solid")

        for idx, column_name in enumerate(export_columns, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = header_font
            cell.alignment = header_alignment
            if column_name in material_set:
                cell.fill = green_fill
            elif column_name in geometry_set:
                cell.fill = yellow_fill
            else:
                cell.fill = blue_fill

            # Match column width to header width only.
            ws.column_dimensions[get_column_letter(idx)].width = max(len(str(column_name)), 1) + 2

        wb.save(export_path)

    def generate_geometry_workbook(self):
        if not self.session.project_dir:
            QMessageBox.warning(self.main_window, "No Project", "Load or create a project first.")
            return

        source_df = getattr(self, "model_dataframe", None)
        if source_df is None:
            source_df = getattr(self.session, "df_mesh", None)

        if source_df is None or source_df.empty:
            QMessageBox.warning(self.main_window, "No Geometry Data", "There is no geometry data to export.")
            return

        export_df = source_df.copy()
        if "Model_name" in export_df.columns:
            ordered_columns = ["Model_name"] + [column for column in export_df.columns if column != "Model_name"]
            export_df = export_df.loc[:, ordered_columns]

        config_dir = Path(self.session.project_dir) / "01_Master_Config"
        export_path = config_dir / "Master_Config.xlsx"

        try:
            if create_master_workbook is None:
                raise RuntimeError("create_master_workbook not available")
            msg = create_master_workbook(str(export_path), df=export_df)
        except FileExistsError as exc:
            QMessageBox.warning(self.main_window, "One Excel Workbook already in 01_Master_Config", str(exc))
            return
        except Exception as exc:  # noqa: BLE001
            QMessageBox.critical(self.main_window, "Export Failed", f"Could not create master workbook:\n{exc}")
            return

        self.update_console(f"[EXPORT] Geometry workbook exported to: {export_path}\n")
        QMessageBox.information(
            self.main_window,
            "Geometry Generated",
            f"Geometry configuration file generated successfully:\n{export_path}",
        )

    def toggle_ccx_solver(self):
        if not self.session.is_solving:
            if self.session.mesh_worker is not None and self.session.mesh_worker.isRunning():
                QMessageBox.information(
                    self.main_window,
                    "Meshing In Progress",
                    "Wait for selected meshing to finish before starting CCX.",
                )
                self._log_pipeline_message("SOLVE|BLOCKED|CCX start blocked because selected meshing is running")
                return

            selected_models = self.get_selected_geometries()
            if not selected_models:
                QMessageBox.warning(self.main_window, "No Selection", "Select at least one row in the analytics table.")
                return
            if not self.session.project_dir:
                QMessageBox.warning(self.main_window, "No Project", "Load or create a project first.")
                return

            setup = self.context.controllers["setup"]
            try:
                if not setup.ensure_configuration_loaded():
                    QMessageBox.warning(
                        self.main_window,
                        "Missing Configuration",
                        "Place exactly one workbook in 01_Master_Config and click Augment or Load.",
                    )
                    return
            except Exception as exc:  # noqa: BLE001
                QMessageBox.warning(self.main_window, "Missing Configuration", f"Could not load workbook:\n{exc}")
                return

            self.session.is_solving = True
            self.update_row_status(selected_models, "Solving...")
            self.ui.btn_toggle_ccx.setText("Abort CCX")
            self.ui.btn_toggle_ccx.setStyleSheet("QPushButton { background-color: #c62828; color: white; font-weight: bold; }")
            self.session.sim_start_time = time.time()
            self.update_sim_timer()
            self.widgets.sim_timer.start(1000)
            self._set_progress_value(0)
            self._log_pipeline_message("SOLVE|QUEUE|Starting selected CCX solve")
            if hasattr(self.ui, "btn_generate_meshes"):
                self.ui.btn_generate_meshes.setEnabled(False)
            if hasattr(self.ui, "btn_tetris"):
                self.ui.btn_tetris.setEnabled(True)

            self.session.ccx_worker = CCXSolverWorker(
                project_dir=self.session.project_dir,
                df_mesh=self.session.df_mesh,
                df_material=self.session.df_material,
                df_step=self.session.df_step,
                target_models=selected_models,
                input_amplitude=self.session.input_amplitude,
            )
            self.session.ccx_worker.progress_updated.connect(self._set_progress_value)
            self.session.ccx_worker.status_updated.connect(self._on_ccx_status_updated)
            self.session.ccx_worker.finished.connect(self._on_ccx_worker_finished)
            self.session.ccx_worker.start()
            return

        if self.session.ccx_worker is not None:
            self.session.ccx_worker.abort()
        self._log_pipeline_message("SOLVE|ABORT|CCX abort requested by user")

        self.widgets.sim_timer.stop()
        self.session.sim_start_time = 0
        self.update_sim_timer()
        self.session.is_solving = False
        self._reset_ccx_toggle_button()
        if hasattr(self.ui, "btn_tetris"):
            self.ui.btn_tetris.setEnabled(False)

        setup = self.context.controllers.get("setup") if hasattr(self.context, "controllers") else None
        if setup is not None and getattr(setup, "tetris_window", None) is not None and setup.tetris_window.isVisible():
            setup.tetris_window.close()

    def load_analytics_data(self):
        if not self.session.project_dir:
            return
        self._refresh_job_table_from_disk()

    def update_plot(self):
        self.widgets.figure.clear()
        self._disconnect_plot_hover()
        source_df = self.session.job_table_df
        if source_df is None or source_df.empty:
            self.widgets.canvas.draw()
            return

        x_col = self.ui.combo_x.currentText() if hasattr(self.ui, "combo_x") else ""
        y_col = self.ui.combo_y.currentText() if hasattr(self.ui, "combo_y") else ""
        z_col = self.ui.combo_z.currentText() if hasattr(self.ui, "combo_z") else ""
        if not x_col or not y_col:
            self.widgets.canvas.draw()
            return

        if x_col not in source_df.columns or y_col not in source_df.columns:
            self.widgets.canvas.draw()
            return

        def _as_series(frame, column_name):
            values = frame.loc[:, column_name]
            if isinstance(values, pd.DataFrame):
                # Duplicate column labels can return a 2D frame; pick first column for plotting.
                values = values.iloc[:, 0]
            return values

        x_series = _as_series(source_df, x_col)
        y_series = _as_series(source_df, y_col)
        plot_df = pd.DataFrame({x_col: x_series, y_col: y_series}, index=source_df.index)
        if "Model_name" in source_df.columns:
            plot_df = plot_df.assign(Model_name=source_df["Model_name"].astype(str))
        else:
            plot_df = plot_df.assign(Model_name=source_df.index.astype(str))
        plot_df[x_col] = pd.to_numeric(plot_df[x_col], errors="coerce")
        plot_df[y_col] = pd.to_numeric(plot_df[y_col], errors="coerce")
        plot_df = plot_df.dropna(subset=[x_col, y_col])

        selected_rows = []
        if hasattr(self.ui, "table_results"):
            selection_model = self.ui.table_results.selectionModel()
            if selection_model is not None:
                selected_rows = sorted({
                    index.row()
                    for index in selection_model.selectedRows()
                    if 0 <= index.row() < len(source_df.index)
                })

        selected_plot_df = pd.DataFrame()
        if selected_rows:
            selected_subset = source_df.iloc[selected_rows].copy()
            selected_x = pd.to_numeric(_as_series(selected_subset, x_col), errors="coerce")
            selected_y = pd.to_numeric(_as_series(selected_subset, y_col), errors="coerce")
            selected_plot_df = pd.DataFrame({x_col: selected_x, y_col: selected_y}, index=selected_subset.index)
            if "Model_name" in selected_subset.columns:
                selected_plot_df = selected_plot_df.assign(Model_name=selected_subset["Model_name"].astype(str))
            else:
                selected_plot_df = selected_plot_df.assign(Model_name=selected_subset.index.astype(str))
            selected_plot_df = selected_plot_df.dropna(subset=[x_col, y_col])
            selected_plot_df = selected_plot_df.sort_values(by=x_col, ascending=True, kind="mergesort")

        if plot_df.empty:
            if hasattr(self.ui, "lbl_info"):
                self.ui.lbl_info.setText(f"No numeric data available for {x_col} vs {y_col}.")
            self.widgets.canvas.draw()
            return

        x_vals = plot_df[x_col]
        y_vals = plot_df[y_col]
        model_vals = plot_df["Model_name"].astype(str)

        if z_col == "--- None (2D) ---":
            ax = self.widgets.figure.add_subplot(111)
            scatter = ax.scatter(x_vals, y_vals, color="tab:blue")
            scatter.set_pickradius(8)
            ax.set_xlabel(x_col)
            ax.set_ylabel(y_col)

            if not selected_plot_df.empty:
                selected_x_vals = selected_plot_df[x_col]
                selected_y_vals = selected_plot_df[y_col]
                selected_models = selected_plot_df["Model_name"].astype(str)
                ax.scatter(selected_x_vals, selected_y_vals, color="darkorange", zorder=5)
                ax.plot(
                    selected_x_vals,
                    selected_y_vals,
                    color="darkorange",
                    linestyle="--",
                    linewidth=2,
                    zorder=4,
                )
                for model_name, x_value, y_value in zip(selected_models, selected_x_vals, selected_y_vals):
                    label_text = f"{model_name}\n({x_value:.2f}, {y_value:.2f})"
                    ax.annotate(
                        label_text,
                        (x_value, y_value),
                        xytext=(5, 5),
                        textcoords="offset points",
                        bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="darkorange", alpha=0.8),
                    )

            hover_points = [
                {"model_name": model_name, "x": x_value, "y": y_value, "z": None}
                for model_name, x_value, y_value in zip(model_vals, x_vals, y_vals)
            ]
            self._bind_plot_hover(scatter, hover_points)
        elif z_col in source_df.columns:
            z_vals = pd.to_numeric(_as_series(source_df.loc[plot_df.index], z_col), errors="coerce")
            plot_df = plot_df.assign(**{z_col: z_vals}).dropna(subset=[z_col])
            if plot_df.empty:
                if hasattr(self.ui, "lbl_info"):
                    self.ui.lbl_info.setText(f"No numeric data available for {x_col}, {y_col}, {z_col}.")
                self.widgets.canvas.draw()
                return
            x_vals = plot_df[x_col]
            y_vals = plot_df[y_col]
            z_vals = plot_df[z_col]
            model_vals = plot_df["Model_name"].astype(str)
            ax = self.widgets.figure.add_subplot(111, projection="3d")
            scatter = ax.scatter(x_vals, y_vals, z_vals, color="tab:blue")
            scatter.set_pickradius(8)
            ax.set_xlabel(x_col)
            ax.set_ylabel(y_col)
            ax.set_zlabel(z_col)

            selected_plot_3d_df = pd.DataFrame()
            if not selected_plot_df.empty:
                selected_plot_3d_df = selected_plot_df.copy()
                selected_subset_3d = source_df.loc[selected_plot_3d_df.index]
                selected_z_vals = pd.to_numeric(_as_series(selected_subset_3d, z_col), errors="coerce")
                selected_plot_3d_df = selected_plot_3d_df.assign(**{z_col: selected_z_vals}).dropna(subset=[z_col])
                selected_plot_3d_df = selected_plot_3d_df.sort_values(by=x_col, ascending=True, kind="mergesort")

            if not selected_plot_3d_df.empty:
                selected_x_vals = selected_plot_3d_df[x_col]
                selected_y_vals = selected_plot_3d_df[y_col]
                selected_z_vals = selected_plot_3d_df[z_col]
                selected_models = selected_plot_3d_df["Model_name"].astype(str)
                ax.scatter(selected_x_vals, selected_y_vals, selected_z_vals, color="darkorange", zorder=5)
                ax.plot(
                    selected_x_vals,
                    selected_y_vals,
                    selected_z_vals,
                    color="darkorange",
                    linestyle="--",
                    linewidth=2,
                    zorder=4,
                )
                for model_name, x_value, y_value, z_value in zip(selected_models, selected_x_vals, selected_y_vals, selected_z_vals):
                    label_text = f"{model_name}\n({x_value:.2f}, {y_value:.2f})"
                    ax.text(
                        x_value,
                        y_value,
                        z_value,
                        label_text,
                        bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="darkorange", alpha=0.8),
                    )

            hover_points = [
                {"model_name": model_name, "x": x_value, "y": y_value, "z": z_value}
                for model_name, x_value, y_value, z_value in zip(model_vals, x_vals, y_vals, z_vals)
            ]
            self._bind_plot_hover(scatter, hover_points)

        try:
            if self.widgets.canvas.geometry().width() > 10 and self.widgets.canvas.geometry().height() > 10:
                self.widgets.figure.tight_layout()
                self.widgets.canvas.draw()
        except Exception:
            pass  # Ignore layout errors when the UI splitter is minimized
